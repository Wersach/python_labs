import csv

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from src.lab04.io_txt_csv import ensure_parent_dir
from src.lib.io_helpers import check_path


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в XLSX (лист Sheet1, первая строка это заголовок).

    Все значения записываются как текст. Ширина колонки равна длине самого
    длинного значения, но не меньше 8. CSV без строк с данными считается пустым.
    TypeError: путь не строка. ValueError: не тот тип файла, пустой CSV,
    строки разной длины. FileNotFoundError: нет входного файла.
    """
    check_path(csv_path, ".csv")
    check_path(xlsx_path, ".xlsx")

    rows = []
    with open(csv_path, encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if len(row) > 0:
                rows.append(row)

    if len(rows) < 2:
        raise ValueError("Пустой CSV или нет строк с данными")
    for row in rows:
        if len(row) != len(rows[0]):
            raise ValueError("строки CSV разной длины")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    for col in range(len(rows[0])):
        width = 8
        for row in rows:
            if len(row[col]) > width:
                width = len(row[col])
        ws.column_dimensions[get_column_letter(col + 1)].width = width

    ensure_parent_dir(xlsx_path)
    wb.save(xlsx_path)


def show_error(func, *args) -> None:
    """Вызывает func(*args) и печатает, какая ошибка получилась."""
    try:
        func(*args)
        print("ошибки нет")
    except (TypeError, ValueError, FileNotFoundError) as e:
        print(type(e).__name__ + ":", e)


if __name__ == "__main__":
    print("CSV -> XLSX")
    for name in ["people", "cities", "students"]:
        src = f"data/samples/{name}.csv"
        dst = f"data/out/{name}.xlsx"
        csv_to_xlsx(src, dst)
        ws = load_workbook(dst)["Sheet1"]
        widths = []
        for col in range(1, ws.max_column + 1):
            widths.append(ws.column_dimensions[get_column_letter(col)].width)
        print(dst, "| лист:", ws.title, "| строк:", ws.max_row, "| колонок:", ws.max_column)
        print("  ширины колонок:", widths)

    print()
    print("Ошибки csv_to_xlsx")
    cases = [
        ("нет файла", "data/samples/net.csv", "data/out/x.xlsx"),
        ("не .csv", "data/samples/people.json", "data/out/x.xlsx"),
        ("не .xlsx", "data/samples/people.csv", "data/out/x.xls"),
        ("пустой файл", "data/samples/bad/empty.csv", "data/out/x.xlsx"),
        ("только заголовок", "data/samples/bad/header_only.csv", "data/out/x.xlsx"),
        ("строки разной длины", "data/samples/bad/short_row.csv", "data/out/x.xlsx"),
    ]
    for label, src, dst in cases:
        print(label + ":", end=" ")
        show_error(csv_to_xlsx, src, dst)
    print("путь не строка:", end=" ")
    show_error(csv_to_xlsx, 5, "data/out/x.xlsx")